from __future__ import annotations
import re
import logging
from pathlib import Path

import httpx
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import settings
from app.models import AgentDocument, Document, DocumentChunk

logger = logging.getLogger(__name__)

_embed_client: httpx.AsyncClient | None = None

def extract_text(path: Path, ext: str) -> str:
    ext = ext.lower()
    if ext == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        parts = []
        for page in reader.pages:
            try:
                parts.append(page.extract_text() or "")
            except Exception:
                logger.exception("Не удалось извлечь текст со страницы PDF")
        return "\n\n".join(p for p in parts if p.strip())

    if ext == "docx":
        import docx

        doc = docx.Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n\n".join(parts)

    raise ValueError(f"Неподдерживаемый формат: .{ext}")



def chunk_text(
    text: str,
    target: int | None = None,
    tolerance: int | None = None,
    overlap: int | None = None,
) -> list[str]:

    target = target or settings.chunk_size
    tolerance = (
        tolerance if tolerance is not None else settings.chunk_tolerance
    )
    overlap = overlap if overlap is not None else settings.chunk_overlap

    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"(?<!\n)\n(?!\n)", " ", text)
    text = text.strip()
    if not text:
        return []

    n = len(text)
    min_size = max(target - tolerance, target // 2)
    max_size = target + tolerance

    overlap = min(overlap, min_size // 2)

    chunks: list[str] = []
    start = 0

    while start < n:
        target_pos = start + target
        min_end = min(start + min_size, n)
        max_end = min(start + max_size, n)

        if n <= max_end:
            end = n
        else:
            window = text[min_end:max_end]
            offset = min_end

            para_positions = []
            p = 0
            while True:
                p = window.find("\n\n", p)
                if p < 0:
                    break
                para_positions.append(offset + p)
                p += 2

            if para_positions:
                best = min(para_positions, key=lambda x: abs(x - target_pos))
                end = best + 2
            else:
                sent_positions = [
                    offset + m.start() + 1
                    for m in re.finditer(r"[.!?](?=\s|$)", window)
                ]
                if sent_positions:
                    best = min(sent_positions, key=lambda x: abs(x - target_pos))
                    end = best
                else:
                    end = start + target

        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)

        if end >= n:
            break

        new_start = end - overlap
        if new_start <= start:
            new_start = start + 1
        start = new_start

    return chunks



def _get_embed_client() -> httpx.AsyncClient:
    global _embed_client
    if _embed_client is None or _embed_client.is_closed:
        _embed_client = httpx.AsyncClient(
            timeout=httpx.Timeout(60.0),
            limits=httpx.Limits(
                max_keepalive_connections=16,
                max_connections=32,
                keepalive_expiry=60.0,
            ),
        )
    return _embed_client


async def close_embedding_client() -> None:
    global _embed_client
    if _embed_client and not _embed_client.is_closed:
        await _embed_client.aclose()


class EmbeddingClient:
    def __init__(self) -> None:
        self.base_url = settings.embedding_base_url.rstrip("/")

    async def _embed(self, inputs: list[str]) -> list[list[float]]:
        if not inputs:
            return []
        inputs = [t if t.strip() else "." for t in inputs]

        results: list[list[float]] = []
        client = _get_embed_client()
        batch_size = 8
        for i in range(0, len(inputs), batch_size):
            batch = inputs[i:i + batch_size]
            response = await client.post(
                f"{self.base_url}/embed",
                json={"inputs": batch, "normalize": True},
            )
            response.raise_for_status()
            results.extend(response.json())
        return results

    async def embed_passages(self, texts: list[str]) -> list[list[float]]:
        return await self._embed([f"passage: {t}" for t in texts])

    async def embed_query(self, text: str) -> list[float]:
        vectors = await self._embed([f"query: {text.strip()}"])
        return vectors[0]


embedding_client = EmbeddingClient()


async def search_chunks(
    db: AsyncSession,
    agent_id: int,
    query: str,
    top_k: int | None = None,
    max_distance: float | None = None,
) -> list[tuple[DocumentChunk, float]]:
    """Возвращает (chunk, distance), отфильтрованные по порогу."""
    top_k = top_k or settings.rag_top_k
    max_distance = (
        max_distance if max_distance is not None
        else settings.rag_max_distance
    )

    query_emb = await embedding_client.embed_query(query)
    distance_expr = DocumentChunk.embedding.cosine_distance(query_emb)

    stmt = (
        select(DocumentChunk, distance_expr.label("distance"))
        .join(Document, Document.id == DocumentChunk.document_id)
        .join(AgentDocument, AgentDocument.document_id == Document.id)
        .where(
            AgentDocument.agent_id == agent_id,
            AgentDocument.is_enabled.is_(True),
            Document.status == "ready",
            DocumentChunk.embedding.is_not(None),
        )
        .order_by(distance_expr)
        .limit(top_k)
    )
    rows = (await db.execute(stmt)).all()
    return [(chunk, float(dist)) for chunk, dist in rows if dist <= max_distance]


async def index_document(
    db: AsyncSession, doc: Document, path: Path, ext: str
) -> None:
    try:
        text = extract_text(path, ext)
        doc.text = text
        doc.status = "processing"
        await db.commit()

        chunks = chunk_text(text)
        if not chunks:
            doc.status = "ready"
            doc.error_message = None
            await db.commit()
            return

        embeddings = await embedding_client.embed_passages(chunks)

        for i, (chunk, emb) in enumerate(zip(chunks, embeddings)):
            db.add(
                DocumentChunk(
                    document_id=doc.id,
                    chunk_index=i,
                    content=chunk,
                    embedding_model=settings.embedding_model,
                    embedding=emb,
                    metadata_={"source": doc.filename},
                )
            )

        doc.status = "ready"
        doc.error_message = None
        await db.commit()

    except Exception as exc:
        logger.exception("Индексация документа %s не удалась", doc.id)
        doc.status = "failed"
        doc.error_message = str(exc)[:1000]
        await db.commit()